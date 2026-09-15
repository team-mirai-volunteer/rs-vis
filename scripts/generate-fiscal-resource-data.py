"""Download official tables and derive compact resource coefficients.

Run: python scripts/generate-fiscal-resource-data.py
Requires requests, openpyxl, numpy and pypdf. Raw files stay under ignored data/.
Money in the IO workbook is BILLION yen; OCCTO power is 10,000 kW.
"""
import hashlib
import json
from pathlib import Path
import re

import numpy as np
import openpyxl
import requests
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / 'data/fiscal-space-sources'
OUT = ROOT / 'app/lib/fiscal-space/data'
CACHE.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)
URLS = {
    'io-2020-108.xlsx': 'https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040187026&fileKind=0',
    'employment-2020-108.xlsx': 'https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040187066&fileKind=0',
    'occto-2026.pdf': 'https://www.occto.or.jp/assets/various/kyoukei/torimatome/260330_kyokyukeikaku_torimatome/260330_kyokei_torimatome.pdf',
}
for name, url in URLS.items():
    if not (CACHE / name).exists():
        response = requests.get(url, timeout=60)
        response.raise_for_status()
        (CACHE / name).write_bytes(response.content)

w = openpyxl.load_workbook(CACHE / 'io-2020-108.xlsx', data_only=True).active
employment = openpyxl.load_workbook(CACHE / 'employment-2020-108.xlsx', data_only=True).active
assert w['DZ1'].value == '(単位 : 10億円)'
codes = [w.cell(i, 1).value for i in range(4, 112)]
names = [w.cell(i, 2).value for i in range(4, 112)]
assert len(codes) == 108 and codes[65] == '461'
jobs = {r[0]: r[2] for r in list(employment.values)[3:]}
# Imputed rent, own-account transport and office supplies have no separate jobs.
assert set(codes) - set(jobs) == {'553', '573', '681'}
workers = np.array([jobs.get(c, 0) for c in codes], dtype=float)
assert workers.sum() == 68707839
production = np.array([w.cell(i, 130).value for i in range(4, 112)]) * 1e9
transactions = np.array([[w.cell(i, j).value or 0 for j in range(3, 111)] for i in range(4, 112)]) * 1e9
domestic_demand = np.array([w.cell(i, 120).value for i in range(4, 112)]) * 1e9
imports = -np.array([w.cell(i, 128).value for i in range(4, 112)]) * 1e9
domestic_share = np.clip(1 - imports / domestic_demand, 0, 1)
# Open IO model: x = [I - (I-M) A]^-1 (I-M) f. No income feedback loop.
matrix = np.eye(108) - domestic_share[:, None] * (transactions / production[None, :])
inverse = np.linalg.inv(matrix)
assert np.max(np.abs(matrix @ inverse - np.eye(108))) < 1e-10

def group(code):
    if code in ['411', '412', '413', '419']: return 'construction'
    if code in ['641', '642', '643', '644']: return 'healthcare'
    if code == '632': return 'research'
    if code in ['321', '329']: return 'electronics'
    if code == '461': return 'electricity'
    return 'general'

groups = [group(c) for c in codes]
sectors = ['general', 'construction', 'healthcare', 'research', 'electronics', 'electricity']
sector_workers = {s: float(sum(workers[i] for i, g in enumerate(groups) if g == s)) for s in sectors}

def column(col):
    a = np.maximum(0, np.array([w.cell(i, col).value or 0 for i in range(4, 112)]))
    return a / a.sum()

def mix(values):
    assert abs(sum(values.values()) - 1) < 1e-10
    return np.array([values.get(c, 0) for c in codes])

profiles = {'household': column(113), 'public-investment': column(116),
    'government': column(114), 'healthcare': mix({'641': 1}), 'childcare': mix({'643': 1}),
    'education': mix({'631': 1}), 'rd': mix({'632': 1}),
    'grid': mix({'419': .6, '331': .3, '669': .1}),
    'generation': mix({'419': .4, '331': .5, '669': .1}),
    'semiconductors': mix({'411': .3, '301': .6, '632': .1}),
    'semiconductor-operation': mix({'321': 1})}
derived = {}
for key, final in profiles.items():
    output = inverse @ (domestic_share * final) * 1e12
    # Small negative recycling coefficients never create extra headroom.
    output = np.maximum(0, output)
    labour = output * workers / production
    derived[key] = {
        'workersPerTrillion': {s: round(float(sum(labour[i] for i, g in enumerate(groups) if g == s)), 6) for s in sectors},
        'electricityYenPerTrillion': round(float(output[codes.index('461')]), 2),
        'domesticOutputPerTrillion': round(float(output.sum()), 2),
    }

# The long-term regional tables use August and (for 3 winter-peaking areas) January.
pdf = PdfReader(CACHE / 'occto-2026.pdf')
summer = pdf.pages[67].extract_text()
winter = pdf.pages[68].extract_text()
regions = ['北海道', '東北', '東京', '中部', '北陸', '関西', '中国', '四国', '九州', '沖縄']
def series(text, name):
    rows = re.findall(r'^' + name + r'\s+([\d, ]+)$', text, re.M)
    numbers = [[int(v.replace(',', '')) / 100 for v in row.split()] for row in rows]
    return [row for row in numbers if len(row) == 10]

power = []
for name in regions:
    values = series(summer, name)
    assert len(values) == 2, (name, values)
    power.append({'region': name, 'season': '8月', 'demandGw': values[0], 'supplyGw': values[1]})
for name in ['北海道', '東北', '北陸']:
    values = series(winter, name)
    assert len(values) == 2, (name, values)
    power.append({'region': name, 'season': '1月', 'demandGw': values[0], 'supplyGw': values[1]})
assert power[2]['demandGw'][0] == 55.01 and power[2]['supplyGw'][0] == 63.3
result = {'ioYear': 2020, 'powerStartYear': 2026, 'checkedAt': '2026-09-15',
    'sources': {name: {'url': url, 'sha256': hashlib.sha256((CACHE / name).read_bytes()).hexdigest()} for name, url in URLS.items()},
    'sectorWorkers': sector_workers, 'profiles': derived, 'power': power}
(OUT / 'resource-reference.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print('Generated resource-reference.json:', len(derived), 'demand profiles,', len(power), 'regional/seasonal power paths')
