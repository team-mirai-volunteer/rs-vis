"""Import IPSS 2023 low fertility, medium mortality annual tables.

Run with Python + openpyxl. Preserve the existing medium series for old scenarios.
"""
import hashlib
import json
import re
import urllib.request
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests/fixtures/fiscal-demographics"
DATA = ROOT / "app/lib/fiscal-space/data"
BASE = "https://www.ipss.go.jp/pp-zenkoku/j/zenkoku2023/db_zenkoku2023/s_tables/"


def main():
    manifest_path = FIXTURES / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    medium = json.loads((DATA / "population-projection.json").read_text(encoding="utf-8"))
    tfr_rows = list(openpyxl.load_workbook(FIXTURES / "ipss2023_ATable1.xlsx", data_only=True).active.values)
    for variant, number, label, column in [("low", 3, "低位", 4)]:
        files = {}
        for suffix in ["8", "9"]:
            name = f"ipss2023_{number}-{suffix}.xlsx"
            url = f"{BASE}{number}-{suffix}.xlsx"
            path = FIXTURES / name
            if not path.exists():
                path.write_bytes(urllib.request.urlopen(url).read())
            raw = path.read_bytes()
            files[name] = {"url": url, "title": f"表{number}-{suffix} 出生{label}・死亡中位", "retrievedOn": date.today().isoformat(), "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw)}
            manifest["files"][name] = files[name]
        population = {}
        for sheet in openpyxl.load_workbook(FIXTURES / f"ipss2023_{number}-9.xlsx", data_only=True).worksheets:
            year = int(re.search(r"\((20\d{2})\)", str(sheet.cell(2, 1).value)).group(1))
            groups = [0] * 18
            ages = set()
            for row in sheet.iter_rows(min_row=5, values_only=True):
                for offset in [0, 5]:
                    age, count = row[offset:offset + 2]
                    if isinstance(age, str) and age.startswith("105"):
                        age = 105
                    if isinstance(age, (int, float)) and isinstance(count, (int, float)) and 0 <= age <= 105:
                        assert age not in ages
                        ages.add(age)
                        groups[min(int(age) // 5, 17)] += round(count * 1000)
            assert len(ages) == 106, (year, ages)
            assert abs(sum(groups) - round(sheet.cell(4, 2).value * 1000)) <= 10
            population[str(year)] = groups
        births = {str(row[1]): round(row[2] * 1000) for row in openpyxl.load_workbook(FIXTURES / f"ipss2023_{number}-8.xlsx", data_only=True).active.values if isinstance(row[1], int) and 2021 <= row[1] <= 2070}
        tfr = {str(row[1]): row[column] for row in tfr_rows if isinstance(row[1], int) and 2020 <= row[1] <= 2070}
        assert len(population) == len(tfr) == 51 and len(births) == 50
        result = {"source": f"国立社会保障・人口問題研究所「日本の将来推計人口（令和5年推計）」出生{label}（死亡中位）推計", "sourcePage": medium["sourcePage"], "variant": f"出生{label}・死亡中位（基本推計）", "retrievedOn": date.today().isoformat(), "units": medium["units"], "referenceYear": 2024, "files": {**files, "ipss2023_ATable1.xlsx": medium["files"]["ipss2023_ATable1.xlsx"]}, "years": medium["years"], "ageGroups": medium["ageGroups"], "population": population, "births": births, "tfr": tfr}
        (DATA / f"population-projection-{variant}.json").write_text(json.dumps(result, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()
