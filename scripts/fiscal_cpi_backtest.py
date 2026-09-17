"""Expanding-window, revised-vintage annual CPI diagnostics (numpy/openpyxl).

Forecasts never use post-origin CPI or gaps. This is a reduced equation check,
not a historical replay of the full fiscal simulator.
"""
import argparse
import csv
import hashlib
import io
import json
import math
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / 'tests/fixtures/fiscal-cpi'
METHODS = ['last-cpi', 'constant-anchor', 'current-held-gap', 'closure-3', 'closure-5', 'closure-10',
           'estimated-gap', 'estimated-anchor-2', 'estimated-intercept']
LABELS = ['直近CPI据置', '基準インフレ固定', '現行係数＋ギャップ据置', '3年で解消', '5年で解消', '10年で解消',
          '解消速度のみ推定', '解消速度・持続性・ギャップ係数推定（基準固定）', '切片も推定']


def load_data():
    book = openpyxl.load_workbook(FIXTURE / '2622gap.xlsx', data_only=True)
    sheet = book['暦年']
    assert sheet.cell(6, 2).value == 'GDP Gap'
    gaps = {int(row[0]): float(row[1]) for row in sheet.iter_rows(min_row=7, values_only=True)}
    raw = list(csv.reader(io.StringIO((FIXTURE / 'cpi-annual-2020.csv').read_bytes().decode('cp932'))))
    assert raw[0][1] == '総合' and raw[2][1] == '0001'
    cpi = {int(r[0].strip()): float(r[1]) for r in raw[6:] if r[0].strip().isdigit() and r[1].strip()}
    years = sorted(set(gaps) & set(cpi))
    assert years == list(range(1994, 2026))
    return [dict(year=y, cpi=cpi[y], gap=gaps[y]) for y in years]


def fit(train, anchor, free_intercept=False):
    """Least squares with 0<=rho<=.99 and beta>=0; solve active boundaries too."""
    prev = np.array([r['cpi'] for r in train[:-1]])
    gap = np.array([r['gap'] for r in train[1:]])
    target = np.array([r['cpi'] for r in train[1:]])
    if free_intercept:
        centers = (prev.mean(), gap.mean(), target.mean())
        x, z, y = prev - centers[0], gap - centers[1], target - centers[2]
    else:
        x, z, y = prev - anchor, gap, target - anchor
    ratio = lambda a, b: float(a / b) if b > 1e-12 else 0.0
    candidates = []
    rho, beta = np.linalg.lstsq(np.column_stack([x, z]), y, rcond=None)[0]
    if 0 <= rho <= .99 and beta >= 0:
        candidates.append((float(rho), float(beta)))
    for rho in [0, .99]:
        candidates.append((rho, max(0.0, ratio(z @ (y - rho * x), z @ z))))
    candidates.append((min(.99, max(0.0, ratio(x @ y, x @ x))), 0.0))
    rho, beta = min(candidates, key=lambda c: float(np.sum((y - c[0] * x - c[1] * z) ** 2)))
    intercept = float(centers[2] - rho * centers[0] - beta * centers[1]) if free_intercept else anchor * (1 - rho)
    return dict(rho=rho, beta=beta, intercept=intercept)


def forecast(train, horizon, anchor=2.0, rho=.25, beta=.05):
    if len(train) < 3 or horizon < 1:
        raise ValueError('Insufficient training observations or invalid horizon')
    if any(b['year'] != a['year'] + 1 for a, b in zip(train, train[1:])):
        raise ValueError('Nonconsecutive training years')
    lag = np.array([r['gap'] for r in train[:-1]])
    now = np.array([r['gap'] for r in train[1:]])
    raw_phi = float(lag @ now / (lag @ lag)) if lag @ lag > 1e-12 else 0.0
    phi = min(1.0, max(0.0, raw_phi))
    fitted = fit(train, anchor)
    free = fit(train, anchor, True)
    coefficients = dict(phi=phi, rawPhi=raw_phi,
                        halfLifeYears=math.log(.5) / math.log(phi) if 0 < phi < 1 else (0 if phi == 0 else None),
                        anchored=fitted, freeIntercept=free)
    paths = {}
    for method in METHODS:
        params = fitted if method == 'estimated-anchor-2' else free if method == 'estimated-intercept' else dict(rho=rho, beta=beta, intercept=anchor * (1 - rho))
        prev = train[-1]['cpi']
        index = 1.0
        path = []
        for h in range(1, horizon + 1):
            factor = max(0, 1 - h / int(method.split('-')[1])) if method.startswith('closure-') else phi ** h if method.startswith('estimated-') else 1
            gap = train[-1]['gap'] * factor
            value = train[-1]['cpi'] if method == 'last-cpi' else anchor if method == 'constant-anchor' else params['intercept'] + params['rho'] * prev + params['beta'] * gap
            index *= 1 + value / 100
            path.append(dict(horizon=h, cpi=value, gap=gap, cumulativePercent=(index - 1) * 100))
            prev = value
        paths[method] = path
    return paths, coefficients


def run(rows, anchor, rho, beta, min_train=10):
    cases, fits = [], []
    for i in range(min_train - 1, len(rows) - 1):
        paths, coefficients = forecast(rows[:i + 1], min(5, len(rows) - i - 1), anchor, rho, beta)
        fits.append(dict(origin=rows[i]['year'], trainingCount=i + 1, **coefficients))
        for h in [1, 3, 5]:
            if i + h >= len(rows):
                continue
            target = rows[i + h]
            actual_cumulative = (math.prod(1 + r['cpi'] / 100 for r in rows[i + 1:i + h + 1]) - 1) * 100
            for method, path in paths.items():
                pred = path[h - 1]
                cases.append(dict(origin=rows[i]['year'], target=target['year'], horizon=h, method=method,
                                  predicted=pred['cpi'], actual=target['cpi'], error=pred['cpi'] - target['cpi'],
                                  predictedGap=pred['gap'], actualGap=target['gap'], gapError=pred['gap'] - target['gap'],
                                  predictedCumulative=pred['cumulativePercent'], actualCumulative=actual_cumulative,
                                  cumulativeError=pred['cumulativePercent'] - actual_cumulative))
    summaries = []
    for h in [1, 3, 5]:
        for method in METHODS:
            group = [r for r in cases if r['horizon'] == h and r['method'] == method]
            if not group:
                continue
            mean = lambda fn: sum(fn(r) for r in group) / len(group)
            summaries.append(dict(horizon=h, method=method, count=len(group), mae=mean(lambda r: abs(r['error'])),
                                  bias=mean(lambda r: r['error']), rmse=math.sqrt(mean(lambda r: r['error'] ** 2)),
                                  gapMae=mean(lambda r: abs(r['gapError'])), cumulativeMae=mean(lambda r: abs(r['cumulativeError']))))
    return dict(cases=cases, summaries=summaries, fits=fits)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--anchor', type=float, required=True)
    parser.add_argument('--rho', type=float, required=True)
    parser.add_argument('--beta', type=float, required=True)
    args = parser.parse_args()
    rows = load_data()
    result = run(rows, args.anchor, args.rho, args.beta)
    periods = []
    for start, end in [(2004, 2019), (2020, 2025)]:
        for h in [1, 3, 5]:
            for method in METHODS:
                group = [r for r in result['cases'] if start <= r['target'] <= end and r['horizon'] == h and r['method'] == method]
                periods.append(dict(start=start, end=end, horizon=h, method=method, count=len(group),
                                    mae=sum(abs(r['error']) for r in group) / len(group),
                                    bias=sum(r['error'] for r in group) / len(group)))
    result['periods'] = periods
    sources = [dict(file=name, url=url, sha256=hashlib.sha256((FIXTURE / name).read_bytes()).hexdigest()) for name, url in [
        ('2622gap.xlsx', 'https://www5.cao.go.jp/keizai3/getsurei/2622gap.xlsx'),
        ('cpi-annual-2020.csv', 'https://www.e-stat.go.jp/stat-search/file-download?fileKind=1&statInfId=000032103940')]]
    result.update(data=rows, sources=sources, config=vars(args), minimumTrainingYears=10,
                  method='expanding-window-revised-vintage-reduced-equation', units='percent / percentage points')
    (ROOT / 'docs/fiscal-space-cpi-backtest.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')
    labels = dict(zip(METHODS, LABELS))
    table = '\n'.join(f"| {r['horizon']}年 | {labels[r['method']]} | {r['count']} | {r['mae']:.3f} | {r['bias']:.3f} | {r['rmse']:.3f} | {r['cumulativeMae']:.3f} |" for r in result['summaries'])
    gap_table = '\n'.join(f"| {r['horizon']}年 | {labels[r['method']]} | {r['gapMae']:.3f} |" for r in result['summaries'] if r['method'] in ['current-held-gap', 'closure-3', 'closure-5', 'closure-10', 'estimated-gap'])
    period_table = '\n'.join(f"| {r['start']}～{r['end']} | {r['horizon']}年 | {labels[r['method']]} | {r['count']} | {r['mae']:.3f} | {r['bias']:.3f} |" for r in periods if r['method'] in ['last-cpi', 'current-held-gap', 'estimated-anchor-2', 'estimated-intercept'])
    fits = '\n'.join(f"| {f['origin']} | {f['phi']:.3f} | {f['halfLifeYears']:.2f} | {f['anchored']['rho']:.3f} | {f['anchored']['beta']:.3f} | {f['freeIntercept']['intercept']:.3f} |" if f['halfLifeYears'] is not None else f"| {f['origin']} | {f['phi']:.3f} | 解消なし | {f['anchored']['rho']:.3f} | {f['anchored']['beta']:.3f} | {f['freeIntercept']['intercept']:.3f} |" for f in result['fits'])
    metric = lambda method, h: next(r for r in result['summaries'] if r['method'] == method and r['horizon'] == h)
    recent = lambda method, h: next(r for r in periods if r['start'] == 2020 and r['method'] == method and r['horizon'] == h)
    report = f'''# CPI持続性・ギャップ解消速度の過去検証

再生成：`npm run backtest:fiscal-cpi`（Python、numpy、openpyxlが必要）。[全結果・入力・推定係数](fiscal-space-cpi-backtest.json)。本体の既定値は変更しない。

## 今回の判断

- ギャップ自体の1年先MAEは据置{metric('current-held-gap', 1)['gapMae']:.3f}ポイントに対し、解消速度推定で{metric('estimated-gap', 1)['gapMae']:.3f}ポイント。解消を入れる意義はある。
- しかしCPIの1年先MAEは据置{metric('current-held-gap', 1)['mae']:.3f}ポイント、解消速度のみ推定{metric('estimated-gap', 1)['mae']:.3f}ポイント。ギャップ経路だけでは改善しない。
- 基準固定で持続性とギャップ係数も推定すると1年先MAEは{metric('estimated-anchor-2', 1)['mae']:.3f}ポイントに縮小。切片も推定すると全期間では改善するが、2020～2025年の1年先は{recent('estimated-intercept', 1)['mae']:.3f}ポイントで現行係数{recent('current-held-gap', 1)['mae']:.3f}ポイントより悪い。3・5年先も同期間では切片自由モデルが悪化している。
- 全期間平均で良い係数をそのまま15年経路に移植しない。次の対象は輸入物価・エネルギー、税率変更等を含む価格式と、近年の環境変化への頑健性。現行係数を正しいと認定した結果でもない。

## データと評価方法

1994～2025暦年。総務省の全国総合CPI・年平均前年比（2020年基準、公表2026-01-23）と内閣府GDPギャップ・暦年（2026-09-14更新）を使う。原本とSHA-256は`tests/fixtures/fiscal-cpi/`と結果JSONに保存。年度・単月前年比とは混ぜない。

- [GDPギャップ原本]({sources[0]['url']})：改定後推計。潜在GDPは観測値ではない。
- [CPI原本]({sources[1]['url']})：長期時系列・全国・年平均・中分類指数前年比、総合列。

最初の10年（1994～2003）を学習し、起点を1年ずつ移動して学習期間を拡張。1年先22件、3年先20件、5年先18件。各起点までの年だけで係数を推定し、その後はCPIもギャップも再帰予測する。将来の実績を入力しない。ただし改定後の過去値を使うため、当時の公表情報だけのリアルタイム予測ではない。起点年の年平均値が利用可能になった時点の診断であり、暦年末に全データが公表済みという意味ではない。

比較式：`CPI(t)=切片+ρ×CPI(t−1)+β×gap(t)`。現行係数の縮約式は切片`基準×(1−ρ)`、基準{args.anchor:.2f}%、ρ={args.rho}、β={args.beta}。基準2%を全過去年に置くのは現行仕様の診断であり、当時の政策目標の再現ではない。エネルギー・政策・能力混雑項は省略しているため、本体全体の予測精度を測った結果ではない。

ギャップは据置、3・5・10年線形解消、`gap(t)=φ×gap(t−1)`の推定を比較。φは切片0の最小二乗を[0,1]へ制約し、生の推定値も保存。ρは[0,0.99]、βは非負の制約付き最小二乗。基準固定の推定と切片自由の推定を分ける。単純比較は直近CPI据置と基準値固定。制約は安定した感度比較のための仮定である。

## CPI誤差（%ポイント）

| 先行期間 | 方法 | 件数 | MAE | バイアス | RMSE | 累積物価のMAE |
| --- | --- | ---: | ---: | ---: | ---: | ---: |
{table}

バイアスは予測−実績。累積は対象期間の`∏(1+CPI/100)−1`。ゼロ・負のインフレを含むためMAPEを使わない。重なる期間の誤差は独立標本ではない。

## 対象年を分けた比較（%ポイント）

| 対象年 | 先行期間 | 方法 | 件数 | MAE | バイアス |
| --- | --- | --- | ---: | ---: | ---: |
{period_table}

2020年以降の別集計はコロナ禍とその後を含む時期の頑健性確認で、政策要因を取り除いた比較ではない。方法や係数をこの集計で再選択・再学習していない。

## ギャップ自体の予測誤差（%ポイント）

| 先行期間 | 方法 | MAE |
| --- | --- | ---: |
{gap_table}

## 起点別の推定

| 起点 | ギャップφ | 半減期（年） | 基準固定ρ | 基準固定β | 切片自由モデルの切片 |
| --- | ---: | ---: | ---: | ---: | ---: |
{fits}

半減期はAR係数からの換算であり、ショックが続く現実の需給ギャップがその年数で必ず解消するという意味ではない。切片・係数の全値はJSON参照。

## 限界と判断基準

消費税率変更、輸入価格、補助金、賃金、期待の変化を未調整。推定係数にはこれらの影響が混ざる。年度別の予測・実績・誤差はJSONで追跡できる。異なる経済環境を含む平均誤差だけで既定値を変更しない。

15年先の予測精度や信頼区間は今回検証していない。まず単純比較を上回るか、ギャップ解消だけで改善するか、切片変更の効果と区別できるかを評価する。採用候補が出ても、エネルギー・制度変更を調整し、期間を分けた検証後に本体への導入を判断する。
'''
    (ROOT / 'docs/fiscal-space-cpi-backtest.md').write_text(report, encoding='utf-8')
    for row in result['summaries']:
        print(row['horizon'], row['method'], round(row['mae'], 3))


if __name__ == '__main__':
    main()
