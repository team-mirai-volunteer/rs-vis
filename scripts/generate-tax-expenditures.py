"""Extract the final MOF summary (not the workbook's draft sheets).

Requires openpyxl. Run from repository root. --download refreshes the source.
"""
import argparse
import hashlib
import json
from pathlib import Path
import unicodedata
import urllib.request

import openpyxl
from openpyxl.utils.datetime import from_excel

URL = 'https://www.mof.go.jp/tax_policy/reference/stm_report/fy2025/houkoku01.xlsx'
SOURCE = Path('data/tax-expenditures/houkoku01-2024.xlsx')
OUTPUT = Path('app/lib/tax-expenditures/data.json')


def clean(value):
    return unicodedata.normalize('NFKC', str(value or '')).strip()


def generate():
    sheet = openpyxl.load_workbook(SOURCE, data_only=True)['総括表']
    starts = [r for r in range(4, 774) if sheet.cell(r, 2).value and not sheet.cell(r-1, 2).value]
    measures = []
    for start, end in zip(starts, starts[1:] + [774]):
        joined = lambda col: ''.join(clean(sheet.cell(r, col).value) for r in range(start, end))
        article, name = joined(2), joined(3)
        # The source merges the name for these two distinct provisions.
        if start == 519:
            name = '関西国際空港用地整備準備金'
        if start == 524:
            name = '中部国際空港整備準備金'
        assert name, start
        deadlines = []
        for r in range(start, end):
            value = sheet.cell(r, 5).value
            if value:
                deadlines.append(from_excel(value).strftime('%Y-%m-%d') if isinstance(value, (int, float)) else clean(value))
        rows, label, block = [], '', 0
        previous_data = False
        for r in range(start, end):
            if clean(sheet.cell(r, 6).value):
                label = clean(sheet.cell(r, 6).value)
            values = [sheet.cell(r, c).value for c in range(7, 16)]
            has_data = any(v is not None and v != '' for v in values)
            if not has_data:
                previous_data = False
                continue
            block = block + 1 if previous_data else 0
            assert block < 3, (start, r)
            previous_data = True
            years = {}
            for year, offset in [(2022, 0), (2023, 3), (2024, 6)]:
                cells = values[offset:offset+3]
                assert all(v in (None, '', '－') or isinstance(v, (int, float)) for v in cells), (r, cells)
                years[str(year)] = [int(v) if isinstance(v, (int, float)) else None for v in cells]
            rows.append({'sourceRow': r, 'section': label, 'entity': ['単体法人', 'うち通算法人（内数）', '連結法人'][block], 'years': years})
        measures.append({'id': f'mof-2024-r{start}', 'article': article, 'name': name,
                         'overview': '\n'.join(clean(sheet.cell(r, 4).value) for r in range(start, end) if sheet.cell(r, 4).value),
                         'deadline': ' '.join(deadlines), 'sourceRow': start, 'rows': rows,
                         'rsProjectId': 127 if start == 196 else None,
                         'aliases': '企業版ふるさと納税 地方創生応援税制' if start == 196 else ''})
    # The middle row is a SUBSET of the first, never add it again.
    # Prior-year columns cover this report's provisions, not every provision
    # covered by the respective historical reports. Validate the current total.
    expected = {2024: 2513286}
    for year, total in expected.items():
        actual = sum(row['years'][str(year)][0] or 0 for m in measures for row in m['rows'] if '内数' not in row['entity'])
        assert actual == total, (year, actual, total)
    extracted = {(row['sourceRow'], col + offset): value
                 for m in measures for row in m['rows']
                 for year, col in [('2022', 7), ('2023', 10), ('2024', 13)]
                 for offset, value in enumerate(row['years'][year]) if value is not None}
    original = {(r, c): sheet.cell(r, c).value for r in range(4, 774) for c in range(7, 16)
                if isinstance(sheet.cell(r, c).value, (int, float))}
    assert extracted == original, 'Numeric source cells were lost or changed'
    result = {'sourceUrl': URL, 'sourceSha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
              'published': '2026-02', 'amountUnit': '千円', 'measures': measures}
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'{len(measures)} source groups; 2024 application total verified; {OUTPUT}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--download', action='store_true')
    args = parser.parse_args()
    if args.download or not SOURCE.exists():
        SOURCE.parent.mkdir(parents=True, exist_ok=True)
        urllib.request.urlretrieve(URL, SOURCE)
    generate()
