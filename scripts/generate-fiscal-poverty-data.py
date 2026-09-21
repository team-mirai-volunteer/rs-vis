"""Download MHLW 2025 CSLC tables and extract 2024 income distributions."""
import hashlib
import json
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / 'tests/fixtures/fiscal-poverty'
BASE = 'https://www.mhlw.go.jp/toukei/saikin/hw/k-tyosa/k-tyosa25/'


def main():
    FIXTURES.mkdir(exist_ok=True)
    files = {}
    for name in ['12', '09']:
        url = BASE + f'xlsx/{name}.xlsx'
        path = FIXTURES / f'mhlw2025-{name}.xlsx'
        if not path.exists():
            path.write_bytes(urllib.request.urlopen(url).read())
        files[path.name] = {'url': url, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'retrievedOn': '2026-09-21'}
    tables = openpyxl.load_workbook(FIXTURES / 'mhlw2025-12.xlsx', data_only=True)
    summary = openpyxl.load_workbook(FIXTURES / 'mhlw2025-09.xlsx', data_only=True)
    bounds = [0, 40, 60, 80, 100, 120, 140, 160, 180, 200, 240, 280, 320, 360, 400, 500, 600, 700, 800, 900, 1000]
    sheet = tables.worksheets[6]
    bins = [{'lower': lower * 10000, 'upper': bounds[i+1] * 10000 if i+1 < len(bounds) else None,
             'allPercent': sheet.cell(9+i, 4).value, 'childPercent': sheet.cell(9+i, 6).value} for i, lower in enumerate(bounds)]
    assert all(isinstance(b['allPercent'], (float, int)) and isinstance(b['childPercent'], (float, int)) for b in bins)
    counts = tables.worksheets[0]
    children = tables.worksheets[2]
    # Table 3's total row is located by its published total household count.
    child_row = next(r for r in children.values if 9174 in r and 1.63 in r)
    rates = summary.worksheets[9]
    result = {'sourceName': '厚生労働省 2025年国民生活基礎調査・統計表1、3、7／概況表11',
              'sourceUrl': BASE + 'dl/06.pdf', 'ratesSourceUrl': BASE + 'dl/03.pdf',
              'incomeYear': 2024, 'surveyYear': 2025, 'checkedAt': '2026-09-21', 'files': files,
              'allPovertyRate': rates.cell(7, 19).value / 100,
              'childPovertyRate': rates.cell(8, 19).value / 100,
              'publishedMedian': rates.cell(13, 19).value * 10000,
              'publishedPovertyLine': rates.cell(14, 19).value * 10000,
              'households': counts.cell(4, 3).value * 1000,
              'meanHouseholdSize': counts.cell(6, 3).value,
              'childHouseholds': counts.cell(4, 6).value * 1000,
              'meanChildren': next(v for v in reversed(child_row) if isinstance(v, (float, int))), 'meanChildHouseholdSize': counts.cell(6, 6).value,
              'bins': bins,
              'notes': ['所得不詳を除く世帯員の分布。公表百分率は丸め値。', '人数は2025年の世帯数×平均人数から近似。所得分布の対象年は2024年。', '世帯構成・税負担との同時分布は未観測。モデルで置く仮定は原統計と区別する。']}
    path = ROOT / 'app/lib/fiscal-space/data/poverty-income-2024.json'
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n', encoding='utf-8', newline='\n')


if __name__ == '__main__':
    main()
