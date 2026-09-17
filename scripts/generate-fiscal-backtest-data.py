"""Freeze annual GFS revenue and GDP from the same 2024 JSNA vintage.

Requires openpyxl. Reads committed workbooks; never silently refreshes sources.
"""
import hashlib
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1] / 'tests/fixtures/fiscal-backtest'
BASE = 'https://www.esri.cao.go.jp/jp/sna/data/data_list/kakuhou/files/2024/'
FILES = ['2024s6_2_jp.xlsx', '2024ffm1n_jp.xlsx']
gfs, gdp = [openpyxl.load_workbook(ROOT / name, data_only=True).worksheets[0] for name in FILES]
assert '１０億円' in gfs.cell(2, 1).value and '１０億円' in gdp.cell(4, 1).value
for row, label in [(9, '11 税'), (44, '12 社会負担'), (77, '143 科料・罰金及び追徴金')]:
    assert gfs.cell(row, 1).value.strip() == label
assert '国内総生産（支出側）' in gdp.cell(48, 1).value
rows = []
for i, year in enumerate(range(1994, 2025)):
    column = 6 + 5 * i
    header = gfs.cell(3, column - 3).value
    assert '年度' in header and int(re.search(r'（(\d{4})）', header).group(1)) == year
    assert gfs.cell(4, column).value == '一般政府'
    assert '年度' in gdp.cell(5, i + 2).value and gdp.cell(7, i + 2).value == year
    values = [gfs.cell(r, column).value for r in [9, 44, 77]] + [gdp.cell(48, i + 2).value]
    assert all(isinstance(v, (int, float)) and v > 0 for v in values)
    tax, social, fines, output = values
    rows.append(dict(year=year, gdpTrillion=output / 1000,
                     taxTrillion=tax / 1000, finesTrillion=fines / 1000, socialTrillion=social / 1000,
                     taxPercentGdp=(tax + fines) / output * 100, socialPercentGdp=social / output * 100,
                     cells=dict(gdp=gdp.cell(48, i + 2).coordinate, tax=gfs.cell(9, column).coordinate,
                                social=gfs.cell(44, column).coordinate, fines=gfs.cell(77, column).coordinate)))
data = dict(source='内閣府 2024年度国民経済計算（2020年基準・2008SNA）、付表6(2) GFS・主要系列表1名目GDP年度',
            url=BASE + '2024_kaku_top.html', vintage='2024年度年次推計・2020年基準', checkedAt='2026-09-17',
            period='fiscal-year', scope='一般政府の税（GFS 11）＋科料・罰金及び追徴金（143）＋社会負担（12）。GDPも年度。IMF暦年系列とは接合しない。',
            sources=[dict(file=name, url=BASE + 'tables/' + name,
                          sha256=hashlib.sha256((ROOT / name).read_bytes()).hexdigest()) for name in FILES], rows=rows)
(ROOT / 'jsna-2024-fiscal-years.json').write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(f'{len(rows)} fiscal years: {rows[0]["year"]}-{rows[-1]["year"]}')

# --- Full macro/fiscal series for the all-series backtest (jsna-2024-fiscal-years-full.json) ---
# The revenue-only fixture above stays byte-identical; this section only adds a second file.
import csv

STOCK_FILE = '2024ss3_jp.xlsx'
CPI_FILE = 'estat-cpi-2020-monthly-japan.csv'
CPI_URL = 'https://www.e-stat.go.jp/stat-search/file-download?statInfId=000032103842&fileKind=1'
stock = openpyxl.load_workbook(ROOT / STOCK_FILE, data_only=True).worksheets[0]
assert '一般政府の部門別資産・負債残高' in stock.cell(1, 1).value and '１０億円' in stock.cell(4, 1).value
assert stock.cell(27, 1).value.strip() == '３．負債'
GFS_ROWS = {'revenue': (8, '1 収入'), 'tax': (9, '11 税'), 'social': (44, '12 社会負担'), 'interestReceived': (66, '1411 利子'),
            'fines': (77, '143 科料・罰金及び追徴金'), 'expenditure': (85, '2 支出'), 'interestPaid': (95, '24 利子'),
            'socialBenefits': (117, '27 社会給付'), 'nonfinancialInvestment': (146, '31 非金融資産の純取得'),
            'netLending': (160, '純貸出(+)／純借入(-)')}
for row, label in GFS_ROWS.values():
    assert gfs.cell(row, 1).value.strip() == label, (row, gfs.cell(row, 1).value)

# CPI: 2020-base national all-items monthly index (e-Stat, Shift_JIS). Fiscal-year average = April..March mean.
with open(ROOT / CPI_FILE, encoding='cp932', newline='') as f:
    cpi_rows = list(csv.reader(f))
assert cpi_rows[0][:2] == ['類・品目', '総合'] and cpi_rows[2][1] == '0001'
monthly = {int(r[0]): float(r[1]) for r in cpi_rows if re.fullmatch(r'\d{6}', r[0] or '')}
assert abs(sum(monthly[202000 + m] for m in range(1, 13)) / 12 - 100) < .05, 'CPI base year must be 2020 = 100'


def fiscal_year_cpi(year):
    months = [year * 100 + m for m in range(4, 13)] + [(year + 1) * 100 + m for m in range(1, 4)]
    return sum(monthly[m] for m in months) / 12


def calendar_year_cpi(year):
    return sum(monthly[year * 100 + m] for m in range(1, 13)) / 12


full_rows = []
for i, year in enumerate(range(1994, 2025)):
    column = 6 + 5 * i
    header = gfs.cell(3, column - 3).value
    assert int(re.search(r'（(\d{4})）', header).group(1)) == year and gfs.cell(4, column).value == '一般政府'
    assert gdp.cell(7, i + 2).value == year
    v = {}
    for key, (row, _) in GFS_ROWS.items():
        cell = gfs.cell(row, column).value
        assert isinstance(cell, (int, float)), (key, year, cell)
        v[key] = cell
    output = gdp.cell(48, i + 2).value
    stock_column = 5 + 4 * i
    stock_header = stock.cell(5, stock_column - 3).value
    assert '暦年末' in stock_header and int(re.search(r'（(\d{4})）', stock_header).group(1)) == year
    assert stock.cell(6, stock_column).value == '合計'
    debt = stock.cell(27, stock_column).value
    assert isinstance(debt, (int, float)) and debt > 0
    # GFS identity (10億円): netLending = revenue − expenditure − nonfinancialInvestment.
    assert abs(v['revenue'] - v['expenditure'] - v['nonfinancialInvestment'] - v['netLending']) < 0.2, year
    # Primary balance = net lending + interest paid − interest received (IMF-style, gross interest both sides).
    primary_balance = v['netLending'] + v['interestPaid'] - v['interestReceived']
    # Primary expenditure = expense − interest paid + net acquisition of nonfinancial assets. This matches the
    # simulator's "total expenditure − interest" definition, where total expenditure includes net investment.
    primary_expenditure = v['expenditure'] - v['interestPaid'] + v['nonfinancialInvestment']
    other_primary_revenue = v['revenue'] - v['tax'] - v['fines'] - v['social'] - v['interestReceived']
    assert abs((v['tax'] + v['fines'] + v['social'] + other_primary_revenue) - primary_expenditure - primary_balance) < 0.2
    cpi_fy, cpi_fy_prev = fiscal_year_cpi(year), fiscal_year_cpi(year - 1)
    full_rows.append(dict(year=year, gdpTrillion=output / 1000,
                          taxTrillion=v['tax'] / 1000, finesTrillion=v['fines'] / 1000, socialTrillion=v['social'] / 1000,
                          revenueTrillion=v['revenue'] / 1000, otherPrimaryRevenueTrillion=other_primary_revenue / 1000,
                          interestReceivedTrillion=v['interestReceived'] / 1000,
                          expenditureTrillion=v['expenditure'] / 1000, interestPaidTrillion=v['interestPaid'] / 1000,
                          socialBenefitsTrillion=v['socialBenefits'] / 1000,
                          nonfinancialInvestmentTrillion=v['nonfinancialInvestment'] / 1000,
                          netLendingTrillion=v['netLending'] / 1000,
                          primaryExpenditureTrillion=primary_expenditure / 1000, primaryBalanceTrillion=primary_balance / 1000,
                          grossDebtTrillion=debt / 1000,
                          cpiFiscalYear=round(cpi_fy, 6), cpiInflation=cpi_fy / cpi_fy_prev - 1,
                          cpiCalendarYear=round(calendar_year_cpi(year), 6),
                          cells=dict(gdp=gdp.cell(48, i + 2).coordinate,
                                     **{key: gfs.cell(row, column).coordinate for key, (row, _) in GFS_ROWS.items()},
                                     grossDebt=stock.cell(27, stock_column).coordinate)))
full = dict(source='内閣府 2024年度国民経済計算（2020年基準・2008SNA）：付表6(2) 一般政府の部門別勘定（GFS）、主要系列表1 名目GDP年度、ストック編付表3 一般政府の部門別資産・負債残高；総務省 消費者物価指数（2020年基準）全国・総合 月次',
            url=BASE + '2024_kaku_top.html', vintage='2024年度年次推計・2020年基準', checkedAt='2026-09-17', period='fiscal-year',
            units='金額は兆円（原表10億円÷1000）。CPIは2020年＝100の指数、cpiInflationは年度平均の前年度比（比率）。',
            definitions=dict(
                revenue='GFS「1 収入」', taxes='GFS「11 税」＋「143 科料・罰金及び追徴金」', socialContributions='GFS「12 社会負担」',
                otherPrimaryRevenue='収入 − 税 − 罰金 − 社会負担 − 受取利子（1411）',
                expenditure='GFS「2 支出」（費用。非金融資産の純取得を含まない）', interestPaid='GFS「24 利子」', interestReceived='GFS「1411 利子」',
                nonfinancialInvestment='GFS「31 非金融資産の純取得」', netLending='GFS「純貸出(+)／純借入(-)」＝収入 − 支出 − 非金融資産の純取得',
                primaryBalance='純貸出 ＋ 支払利子 − 受取利子（IMF方式。受払利子ともグロス）',
                primaryExpenditure='支出 − 支払利子 ＋ 非金融資産の純取得（シミュレータの「総支出−支払利子」と同じ範囲）',
                grossDebt='ストック編付表3「３．負債」一般政府合計。暦年末・時価評価・SNAの全負債（債務証券・借入・その他の負債等）。年度末（3月末）ではなく、IMF表4の額面・連結ベース総債務とも定義が異なる',
                cpiFiscalYear='当年4月〜翌年3月の月次指数の単純平均（2020年＝100）', cpiCalendarYear='暦年平均（参考）'),
            caveats=['負債残高は暦年末（12月末）値を年度（翌3月末）の代理として使う。年度末値はこの表にない。',
                     '負債はSNA時価評価の全負債で、IMF・国債統計の額面ベース債務より水準が高い。比率の水準ではなく変化の検証に使う。',
                     'CPIは月次から年度平均を計算しているため、暦年平均を用いる既存 fixture（tests/fixtures/fiscal-cpi/cpi-annual-2020.csv）とは値が異なる。'],
            sources=[dict(file=name, url=BASE + 'tables/' + name, sha256=hashlib.sha256((ROOT / name).read_bytes()).hexdigest())
                     for name in FILES + [STOCK_FILE]]
                    + [dict(file=CPI_FILE, url=CPI_URL, sha256=hashlib.sha256((ROOT / CPI_FILE).read_bytes()).hexdigest(),
                            note='e-Stat 消費者物価指数 2020年基準 全国 中分類指数 月次（1970年1月〜）。Shift_JIS。総務省統計局。')],
            rows=full_rows)
(ROOT / 'jsna-2024-fiscal-years-full.json').write_text(json.dumps(full, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(f'full series: {len(full_rows)} fiscal years, debt {full_rows[-1]["grossDebtTrillion"]:.1f} trillion at end-2024, '
      f'FY2024 CPI {full_rows[-1]["cpiFiscalYear"]:.2f} ({full_rows[-1]["cpiInflation"] * 100:.2f}%)')
